"""
TK 平台账单对账脚本
用法:
  本地文件:    python reconcile.py /path/to/TK3631.xlsx
  钉钉URL:     python reconcile.py --url "<dingtalk_download_url>"
  钉钉base64:  python reconcile.py --base64 "<base64_encoded_content>"
  具名参数:    python reconcile.py --excel /path/to/TK3631.xlsx
  指定配置:    python reconcile.py TK3631.xlsx --env /path/to/.env
"""
import os
import sys
import argparse
import base64
import tempfile
import urllib.request
import pg8000
import openpyxl
from collections import defaultdict
import calendar
from datetime import datetime


# ──────────────────────────────────────────────
# 工具函数
# ──────────────────────────────────────────────

def load_env(env_path):
    """从 .env 文件读取 KEY=VALUE 配置"""
    config = {}
    if not os.path.exists(env_path):
        raise FileNotFoundError("配置文件不存在：%s" % env_path)
    with open(env_path, 'r', encoding='utf-8') as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith('#'):
                continue
            if '=' in line:
                key, _, val = line.partition('=')
                config[key.strip()] = val.strip()
    return config


def db_connect(config, database):
    """创建数据库连接"""
    return pg8000.connect(
        host=config['DB_HOST'],
        port=int(config['DB_PORT']),
        database=database,
        user=config['DB_USER'],
        password=config['DB_PASS']
    )


# ──────────────────────────────────────────────
# 第一步：读取 Excel
# ──────────────────────────────────────────────

def read_excel(excel_path):
    print("\n[1/4] 读取 Excel 数据：%s" % excel_path)
    wb = openpyxl.load_workbook(excel_path)
    ws = wb['Order details']
    rows = list(ws.iter_rows(values_only=True))
    headers = rows[0]

    related_idx   = next(i for i, h in enumerate(headers) if h and 'Related order ID' in h)
    settle_idx    = next(i for i, h in enumerate(headers) if h and h.strip() == 'Total settlement amount')
    settled_t_idx = next(i for i, h in enumerate(headers) if h and h.strip() == 'Order settled time')

    excel_data = defaultdict(lambda: {'amount': 0, 'settled_time': None})
    for r in rows[1:]:
        related = r[related_idx]
        if related is None:
            continue
        s = str(related).strip()
        if not s or s.startswith('/'):
            continue
        try:
            amount = int(r[settle_idx]) if r[settle_idx] is not None else 0
        except Exception:
            amount = 0
        t = str(r[settled_t_idx]).strip() if r[settled_t_idx] else None
        excel_data[s]['amount'] += amount
        if t and (excel_data[s]['settled_time'] is None or t > excel_data[s]['settled_time']):
            excel_data[s]['settled_time'] = t

    # 计算账期 start_time / end_time
    all_dates  = [v['settled_time'] for v in excel_data.values() if v['settled_time']]
    min_date   = min(all_dates)
    dt_ref     = datetime.strptime(min_date, '%Y/%m/%d')
    last_day   = calendar.monthrange(dt_ref.year, dt_ref.month)[1]
    start_time = '%d-%02d-01 00:00:00' % (dt_ref.year, dt_ref.month)
    end_time   = '%d-%02d-%02d 23:59:59' % (dt_ref.year, dt_ref.month, last_day)

    total_rows   = len(excel_data)
    total_amount = sum(v['amount'] for v in excel_data.values())

    print("  Excel 唯一订单数：%d" % total_rows)
    print("  start_time  ：%s" % start_time)
    print("  end_time    ：%s" % end_time)
    print("  Excel 总结算金额：%d" % total_amount)

    return excel_data, start_time, end_time, total_rows, total_amount


# ──────────────────────────────────────────────
# 第二步：从 eccang 获取店铺信息
# ──────────────────────────────────────────────

def get_shop_info(config, excel_data):
    print("\n[2/4] 连接 eccang 获取店铺信息...")
    conn = db_connect(config, 'eccang')
    cur  = conn.cursor()
    company_code = user_account = platform = None
    for oid in list(excel_data.keys()):
        cur.execute(
            "SELECT company_code, user_account, platform FROM erp_eb.orders WHERE refrence_no=%s LIMIT 1",
            (oid,)
        )
        row = cur.fetchone()
        if row and row[0]:
            company_code = row[0]
            user_account = row[1]
            platform     = row[2]
            print("  使用 orderId：%s" % oid)
            break
    cur.close()
    conn.close()

    if not company_code:
        print("  ERROR: 未能从数据库获取店铺信息，请检查数据")
        sys.exit(1)

    print("  company_code：%s" % company_code)
    print("  user_account：%s" % user_account)
    print("  platform    ：%s" % platform)
    return company_code, user_account, platform


# ──────────────────────────────────────────────
# 第三步：从 bi 查询结算明细
# ──────────────────────────────────────────────

def query_bi_settlement(config, company_code, user_account, platform, start_time, end_time):
    print("\n[3/4] 连接 bi 查询结算明细...")
    conn = db_connect(config, 'bi')
    cur  = conn.cursor()
    sql = """
SELECT order_id, SUM(fee_amount) AS settlementAmount, account_check_result
FROM bi_xpj_profit_settlement.xpj_ads_profit_settlement_order_detail
WHERE 1=1
  AND platform = %s
  AND company_code = %s
  AND user_account = %s
  AND account_check_result IN ('\u5339\u914d\u5230\u8ba2\u5355', '\u8ba2\u5355\u5df2\u4f5c\u5e9f', '\u5bf9\u8d26\u6210\u529f')
  AND statement_time >= %s
  AND statement_time <= %s
  AND source_type = 1
GROUP BY order_id, account_check_result
ORDER BY settlementAmount DESC
"""
    cur.execute(sql, (platform, company_code, user_account, start_time, end_time))
    db_rows = cur.fetchall()
    cur.close()
    conn.close()

    db_data = defaultdict(int)
    for r in db_rows:
        order_id = str(r[0]).strip() if r[0] else ''
        amount   = int(r[1]) if r[1] is not None else 0
        if order_id:
            db_data[order_id] += amount

    print("  数据库唯一订单数：%d" % len(db_data))
    print("  数据库总结算金额：%d" % sum(db_data.values()))
    return db_data


# ──────────────────────────────────────────────
# 第四步：逐条对比
# ──────────────────────────────────────────────

def compare(excel_data, db_data):
    print("\n[4/4] 开始对比...")
    diff_records = []
    auto_matched = []

    for oid, v in excel_data.items():
        excel_amt = v['amount']
        if oid in db_data:
            db_amt = db_data[oid]
            if excel_amt != db_amt:
                diff_records.append({
                    'orderId':      oid,
                    'excel_amount': excel_amt,
                    'db_amount':    db_amt,
                    'diff':         excel_amt - db_amt,
                    'type':         '\u91d1\u989d\u5dee\u5f02'
                })
        else:
            if excel_amt == 0:
                db_data[oid] = 0
                auto_matched.append(oid)
            else:
                diff_records.append({
                    'orderId':      oid,
                    'excel_amount': excel_amt,
                    'db_amount':    None,
                    'diff':         None,
                    'type':         'Excel\u6709/DB\u65e0'
                })

    for oid, db_amt in db_data.items():
        if oid not in excel_data:
            diff_records.append({
                'orderId':      oid,
                'excel_amount': None,
                'db_amount':    db_amt,
                'diff':         None,
                'type':         'DB\u6709/Excel\u65e0'
            })

    db_total_rows   = len(db_data)
    db_total_amount = sum(db_data.values())
    print("  自动补齐（Excel=0/DB无）：%d 条" % len(auto_matched))
    return diff_records, auto_matched, db_total_rows, db_total_amount


# ──────────────────────────────────────────────
# 打印结果摘要
# ──────────────────────────────────────────────

def print_summary(company_code, user_account, platform, start_time, end_time,
                  excel_total_rows, db_total_rows,
                  excel_total_amount, db_total_amount,
                  auto_matched, diff_records, is_matched):
    print("\n" + "=" * 60)
    if is_matched:
        print("  \U0001f389\U0001f38a 对账结果：完全匹配！恭喜恭喜！太棒了！ \U0001f38a\U0001f389")
        print("=" * 60)
        print("  客户代码    ：%s" % company_code)
        print("  店    铺    ：%s" % user_account)
        print("  平    台    ：%s" % platform)
        print("  账单开始时间：%s" % start_time)
        print("  账单结束时间：%s" % end_time)
        print("  Excel 订单总数    ：%d" % excel_total_rows)
        print("  数据库订单总数    ：%d" % db_total_rows)
        print("  Excel 总结算金额  ：%d" % excel_total_amount)
        print("  数据库总结算金额  ：%d" % db_total_amount)
        print("\n  \u2705 所有订单金额完美匹配，账单核对无误！")
        print("  \U0001f4aa 你的数据真的很干净，继续保持！加油！")
    else:
        print("  \U0001f622\U0001f614 对账结果：存在差异，需要核查...")
        print("=" * 60)
        print("  客户代码    ：%s" % company_code)
        print("  店    铺    ：%s" % user_account)
        print("  平    台    ：%s" % platform)
        print("  账单开始时间：%s" % start_time)
        print("  账单结束时间：%s" % end_time)
        print("  Excel 订单总数    ：%d" % excel_total_rows)
        print("  数据库订单总数    ：%d" % db_total_rows)
        print("  Excel 总结算金额  ：%d" % excel_total_amount)
        print("  数据库总结算金额  ：%d" % db_total_amount)
        print("  自动补齐订单数    ：%d" % len(auto_matched))
        print("  差异条数          ：%d" % len(diff_records))
        print("\n  \U0001f614 发现差异，别灰心，一起查清楚！加油！\U0001f4aa")
    print("=" * 60)


# ──────────────────────────────────────────────
# 输出 Excel 报告
# ──────────────────────────────────────────────

def write_excel_report(out_path, company_code, user_account, platform,
                       start_time, end_time,
                       excel_total_rows, db_total_rows,
                       excel_total_amount, db_total_amount,
                       auto_matched, diff_records,
                       excel_data, db_data, is_matched):

    wb_out = openpyxl.Workbook()

    # —— 对账摘要 Sheet ——
    ws_sum = wb_out.active
    ws_sum.title = '\u5bf9\u8d26\u6458\u8981'
    ws_sum.append(['\u5bf9\u8d26\u6458\u8981'])
    ws_sum.append([])
    ws_sum.append(['start_time', start_time])
    ws_sum.append(['end_time',   end_time])
    ws_sum.append([])
    ws_sum.append(['\u5ba2\u6237\u4ee3\u7801',   company_code])
    ws_sum.append(['\u5e97\u94fa',               user_account])
    ws_sum.append(['\u5e73\u53f0',               platform])
    ws_sum.append(['\u8d26\u5355\u5f00\u59cb\u65f6\u95f4', start_time])
    ws_sum.append(['\u8d26\u5355\u7ed3\u675f\u65f6\u95f4', end_time])
    ws_sum.append([])
    ws_sum.append(['', 'Excel', '\u6570\u636e\u5e93', '\u662f\u5426\u4e00\u81f4'])
    ws_sum.append(['\u8ba2\u5355\u603b\u6570',   excel_total_rows,   db_total_rows,
                   '\u2705' if excel_total_rows == db_total_rows else '\u274c'])
    ws_sum.append(['\u603b\u7ed3\u7b97\u91d1\u989d', excel_total_amount, db_total_amount,
                   '\u2705' if excel_total_amount == db_total_amount else '\u274c'])
    ws_sum.append(['\u81ea\u52a8\u8865\u9f50\u8ba2\u5355\u6570', len(auto_matched), '', ''])
    ws_sum.append(['\u5dee\u5f02\u6761\u6570', len(diff_records), '',
                   '\u2705' if len(diff_records) == 0 else '\u274c'])
    ws_sum.append([])
    if is_matched:
        ws_sum.append(['\u5bf9\u8d26\u7ed3\u8bba', '\U0001f389 \u5b8c\u5168\u5339\u914d\uff0c\u8d26\u5355\u6838\u5bf9\u65e0\u8bef\uff01\u592a\u68d2\u4e86\uff01'])
    else:
        ws_sum.append(['\u5bf9\u8d26\u7ed3\u8bba', '\U0001f622 \u5b58\u5728\u5dee\u5f02\uff0c\u8be6\u89c1\u3010\u5dee\u5f02\u660e\u7ec6\u3011Sheet'])
    for col, w in zip('ABCD', [20, 24, 24, 12]):
        ws_sum.column_dimensions[col].width = w

    # —— 差异明细 Sheet ——
    ws_diff = wb_out.create_sheet('\u5dee\u5f02\u660e\u7ec6')
    ws_diff.append(['orderId', 'Excel\u91d1\u989d', '\u6570\u636e\u5e93\u91d1\u989d', '\u5dee\u5024', '\u5dee\u5f02\u7c7b\u578b'])
    for d in diff_records:
        ws_diff.append([
            d['orderId'],
            d['excel_amount'] if d['excel_amount'] is not None else '',
            d['db_amount']    if d['db_amount']    is not None else '',
            d['diff']         if d['diff']         is not None else '',
            d['type']
        ])
    for col, w in zip('ABCDE', [25, 16, 16, 14, 16]):
        ws_diff.column_dimensions[col].width = w

    # —— Excel 明细 Sheet ——
    ws_excel = wb_out.create_sheet('Excel\u660e\u7ec6')
    ws_excel.append(['orderId', 'settlementAmount', 'order_settled_time'])
    excel_list = [
        (oid, v['amount'],
         datetime.strptime(v['settled_time'], '%Y/%m/%d').strftime('%Y-%m-%d') if v['settled_time'] else '')
        for oid, v in excel_data.items()
    ]
    excel_list.sort(key=lambda x: x[1], reverse=True)
    for r in excel_list:
        ws_excel.append(list(r))
    for col, w in zip('ABC', [25, 20, 20]):
        ws_excel.column_dimensions[col].width = w

    # —— 数据库明细 Sheet ——
    ws_db = wb_out.create_sheet('\u6570\u636e\u5e93\u660e\u7ec6')
    ws_db.append(['orderId', 'settlementAmount'])
    for r in sorted(db_data.items(), key=lambda x: x[1], reverse=True):
        ws_db.append(list(r))
    for col, w in zip('AB', [25, 20]):
        ws_db.column_dimensions[col].width = w

    wb_out.save(out_path)
    print("\n\U0001f4c4 对账结果已输出：%s" % out_path)
    print("   Sheet：对账摘要 / 差异明细 / Excel明细 / 数据库明细")


# ──────────────────────────────────────────────
# 主入口
# ──────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(description='TK 平台账单对账工具')
    # 支持位置参数：python reconcile.py TK3631.xlsx
    # 同时兼容具名参数：python reconcile.py --excel TK3631.xlsx
    parser.add_argument('excel_pos', nargs='?', default=None, metavar='EXCEL_FILE',
                        help='Excel 文件路径（位置参数）')
    parser.add_argument('--excel', default=None, help='Excel 文件路径（具名参数）')
    parser.add_argument('--url',    default=None, help='钉钉文件下载 URL，脚本自动下载到临时目录')
    parser.add_argument('--base64', default=None, dest='b64',
                        help='钉钉对话框文件的 base64 编码内容，脚本自动解码到临时目录')
    parser.add_argument('--env',   default=os.path.join(os.path.dirname(__file__), '..', 'assets', '.env'),
                        help='.env 配置文件路径（默认：assets/.env）')
    parser.add_argument('--out',   default=None, help='输出 Excel 路径（默认：与输入同目录）')
    args = parser.parse_args()

    _tmp_file = None  # 记录临时文件，用于清理

    if args.url:
        # ── 情况A1：从 URL 下载 ──
        print("  检测到文件 URL，正在下载...")
        _tmp = tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx')
        _tmp.close()
        _tmp_file = _tmp.name
        try:
            urllib.request.urlretrieve(args.url, _tmp_file)
            print("  文件已下载至：%s" % _tmp_file)
        except Exception as e:
            print("  ERROR: 文件下载失败 —— %s" % e)
            print("  请确认 URL 是否有效，或改用本地路径：python reconcile.py /path/to/file.xlsx")
            sys.exit(1)
        excel_file = _tmp_file

    elif args.b64:
        # ── 情况A2：从 base64 内容解码 ──
        print("  检测到 base64 文件内容，正在解码...")
        _tmp = tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx')
        _tmp.close()
        _tmp_file = _tmp.name
        try:
            raw = args.b64.strip()
            file_bytes = base64.b64decode(raw)
            with open(_tmp_file, 'wb') as f:
                f.write(file_bytes)
            print("  文件已解码至：%s（%d bytes）" % (_tmp_file, len(file_bytes)))
        except Exception as e:
            print("  ERROR: base64 解码失败 —— %s" % e)
            print("  请确认 base64 内容完整无误")
            sys.exit(1)
        excel_file = _tmp_file

    else:
        # ── 情况B：本地文件路径 ──
        excel_file = args.excel_pos or args.excel
        if not excel_file:
            parser.error("请提供 Excel 文件路径或附件，例如：\n"
                         "  python reconcile.py TK3631.xlsx\n"
                         "  python reconcile.py --url <钉钉文件链接>\n"
                         "  python reconcile.py --base64 <base64内容>")

    # 输出路径
    if args.out:
        out_path = args.out
    else:
        base, _ = os.path.splitext(excel_file)
        out_path = base + '_reconcile_result.xlsx'

    print("=" * 60)
    print("  TK 平台账单对账脚本")
    print("=" * 60)

    # 加载配置
    env_path = os.path.abspath(args.env)
    config   = load_env(env_path)
    print("\n  配置文件：%s" % env_path)

    # 执行四步流程
    excel_data, start_time, end_time, excel_total_rows, excel_total_amount = read_excel(excel_file)
    company_code, user_account, platform = get_shop_info(config, excel_data)
    db_data = query_bi_settlement(config, company_code, user_account, platform, start_time, end_time)
    diff_records, auto_matched, db_total_rows, db_total_amount = compare(excel_data, db_data)

    is_matched = (
        excel_total_rows == db_total_rows and
        excel_total_amount == db_total_amount and
        len(diff_records) == 0
    )

    # 打印摘要
    print_summary(company_code, user_account, platform, start_time, end_time,
                  excel_total_rows, db_total_rows,
                  excel_total_amount, db_total_amount,
                  auto_matched, diff_records, is_matched)

    # 写出 Excel 报告
    write_excel_report(out_path, company_code, user_account, platform,
                       start_time, end_time,
                       excel_total_rows, db_total_rows,
                       excel_total_amount, db_total_amount,
                       auto_matched, diff_records,
                       excel_data, db_data, is_matched)

    # 询问是否发送摘要
    print("\n" + "-" * 60)
    try:
        ans = input("是否将对账结果摘要发送到对话框？(y/n): ").strip().lower()
    except (EOFError, KeyboardInterrupt):
        ans = 'n'

    if ans == 'y':
        print("\n" + "=" * 60)
        print("  \U0001f4cb 对账结果摘要")
        print("  客户代码    ：%s" % company_code)
        print("  店    铺    ：%s" % user_account)
        print("  平    台    ：%s" % platform)
        print("  账单开始时间：%s" % start_time)
        print("  账单结束时间：%s" % end_time)
        print("  Excel 订单总数  ：%d   数据库订单总数  ：%d" % (excel_total_rows, db_total_rows))
        print("  Excel 总金额    ：%d   数据库总金额    ：%d" % (excel_total_amount, db_total_amount))
        if is_matched:
            print("  结论：\U0001f389 完全匹配！")
        else:
            print("  结论：\U0001f622 存在 %d 条差异，详见输出文件" % len(diff_records))
        print("=" * 60)

    # 清理临时下载文件
    if _tmp_file and os.path.exists(_tmp_file):
        try:
            os.remove(_tmp_file)
        except Exception:
            pass


if __name__ == '__main__':
    main()
